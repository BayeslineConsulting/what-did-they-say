#!/usr/bin/env python3
"""Recover master_codebook.json and history.json from a prior run's workbook.
Usage: python state_from_workbook.py prior.xlsx OUT_DIR"""
import json,sys,os
from openpyxl import load_workbook
wb=load_workbook(sys.argv[1],data_only=True);out=sys.argv[2];os.makedirs(out,exist_ok=True)
hist=json.loads(wb['History_JSON']['A1'].value)
codes={};cats={}
for r in wb['Codebook'].iter_rows(min_row=2,values_only=True):
    if not r[0]: continue
    codes[r[1]]=dict(code=r[0],method=r[2],definition=r[3],include=r[4] or '',exclude=r[5] or '',category=r[6],added_run=r[9],last_seen_run=r[10])
for r in wb['Categories'].iter_rows(min_row=2,values_only=True):
    if r[0]: cats[r[0]]=dict(theory_construct=r[4] or '',theory_source=r[5] or '',justification=r[6] or '')
json.dump(dict(version=hist[-1]['conference'],codes=codes,categories=cats),open(os.path.join(out,'master_codebook.json'),'w'),indent=1,ensure_ascii=False)
json.dump(hist,open(os.path.join(out,'history.json'),'w'),indent=1,ensure_ascii=False)
print(f"recovered {len(codes)} codes, {len(cats)} categories, {len(hist)} history runs")
